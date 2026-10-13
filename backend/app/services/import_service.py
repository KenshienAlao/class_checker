"""Student Import service: parse .xlsx workbooks and persist rows atomically."""

import io
from uuid import UUID

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.student import Student
from app.repositories.classroom_repository import ClassroomRepository
from app.repositories.student_repository import StudentRepository
from app.schemas.import_schema import ImportRejection, ImportSummary

_MAX_FILE_SIZE = 5 * 1024 * 1024
_MAX_DATA_ROWS = 1000
_MAX_STUDENT_ID = 64
_MAX_FULL_NAME = 255
_MAX_PROGRAM = 255

_DUPLICATE_CONSTRAINT_NAME = "uq_students_classroom_id_student_id"


def _is_duplicate_student_id_constraint(exc: IntegrityError) -> bool:
    """Identify only the classroom-scoped student-ID unique constraint."""
    orig = getattr(exc, "orig", None)
    diag = getattr(orig, "diag", None)
    constraint_name = getattr(diag, "constraint_name", None)
    if constraint_name == _DUPLICATE_CONSTRAINT_NAME:
        return True
    text = str(orig if orig is not None else exc).lower()
    return _DUPLICATE_CONSTRAINT_NAME in text


def _str_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


class ImportService:
    def __init__(
        self,
        db: Session,
        student_repository: StudentRepository,
        classroom_repository: ClassroomRepository,
    ) -> None:
        self.db = db
        self.student_repository = student_repository
        self.classroom_repository = classroom_repository

    def import_file(self, classroom_id: int, owner_id: UUID, file: UploadFile) -> ImportSummary:
        if self.classroom_repository.get_owned(classroom_id, owner_id) is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Classroom not found")

        contents = file.file.read()
        if len(contents) > _MAX_FILE_SIZE:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="File exceeds 5 MB")
        filename = (file.filename or "").lower()
        if not filename.endswith(".xlsx"):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Only .xlsx files are accepted")

        try:
            wb = load_workbook(filename=io.BytesIO(contents), data_only=True, read_only=True)
        except Exception:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Invalid workbook") from None

        ws = wb.active
        if ws is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Invalid workbook")
        try:
            rows = list(ws.iter_rows(values_only=True))
        except Exception:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Invalid workbook") from None

        if not rows:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Missing required headers")

        header_cells = rows[0]
        header_map: dict[str, int] = {}
        for idx, h in enumerate(header_cells):
            name = _str_cell(h)
            if name:
                header_map[name.lower()] = idx

        # case-insensitive trimmed matching
        normalized = {
            k.strip().lower() if isinstance(k, str) else str(k).strip().lower(): v
            for k, v in header_map.items()
        }
        header_map = normalized

        si_idx = header_map.get("student id")
        fn_idx = header_map.get("full name")
        prog_idx = header_map.get("program")
        if si_idx is None or fn_idx is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Missing required headers")

        data_rows = rows[1:]
        if len(data_rows) > _MAX_DATA_ROWS:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Too many rows")

        # validate + duplicates
        validation_rejected: list[ImportRejection] = []
        duplicate_rejected: list[ImportRejection] = []
        parsed: list[tuple[str, str, str | None]] = []
        seen: set[str] = set()
        existing_ids = set(self.student_repository.list_existing_student_ids(classroom_id))

        for i, row in enumerate(data_rows, start=1):
            sid = _str_cell(row[si_idx] if si_idx < len(row) else None)
            fname = _str_cell(row[fn_idx] if fn_idx < len(row) else None)
            prog_raw = _str_cell(
                row[prog_idx] if (prog_idx is not None and prog_idx < len(row)) else None
            )
            prog = prog_raw if prog_raw != "" else None

            is_duplicate = False
            errs = []
            if not sid:
                errs.append("student_id is required")
            elif len(sid) > _MAX_STUDENT_ID:
                errs.append("student_id too long")
            if not fname:
                errs.append("full_name is required")
            elif len(fname) > _MAX_FULL_NAME:
                errs.append("full_name too long")
            if prog is not None and len(prog) > _MAX_PROGRAM:
                errs.append("program too long")
            if sid:
                if sid in seen:
                    is_duplicate = True
                if sid in existing_ids:
                    is_duplicate = True

            if errs:
                validation_rejected.append(ImportRejection(row=i, error="; ".join(errs)))
                continue
            if is_duplicate:
                duplicate_rejected.append(ImportRejection(row=i, error="duplicate student_id"))
                continue
            seen.add(sid)
            parsed.append((sid, fname, prog))

        if validation_rejected:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail=ImportSummary(
                    created=0, skipped=0, rejected=validation_rejected
                ).model_dump(),
            )
        if duplicate_rejected:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                detail=ImportSummary(
                    created=0, skipped=0, rejected=duplicate_rejected
                ).model_dump(),
            )

        try:
            for sid, fname, prog in parsed:
                self.student_repository.db.add(
                    Student(
                        classroom_id=classroom_id, student_id=sid, full_name=fname, program=prog
                    )
                )
            self.student_repository.db.commit()
        except IntegrityError as exc:
            self.student_repository.db.rollback()
            if not _is_duplicate_student_id_constraint(exc):
                raise
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                detail="Duplicate student ID",
            ) from None

        return ImportSummary(created=len(parsed), skipped=0, rejected=[])
