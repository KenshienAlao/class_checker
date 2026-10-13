"""Content-level integration tests for the XLSX export service."""

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.classroom import Classroom
from app.models.student import Student
from app.models.user import User
from app.repositories.attendance_record_repository import AttendanceRecordRepository
from app.repositories.attendance_session_repository import AttendanceSessionRepository
from app.repositories.classroom_repository import ClassroomRepository
from app.repositories.student_repository import StudentRepository
from app.services.attendance_history_service import (
    AttendanceHistoryGrid,
    AttendanceHistoryService,
)
from app.services.attendance_service import AttendanceService, SessionReview
from app.services.excel_service import ExcelService

_EXPORTED_AT = datetime(2026, 3, 2, 9, 0, tzinfo=UTC)


def _owner(db: Session) -> User:
    user = User(id=uuid4(), name="Owner", email=f"{uuid4()}@example.com", password="n/a")
    db.add(user)
    db.flush()
    return user


def _classroom(db: Session, owner: User) -> Classroom:
    return ClassroomRepository(db).create(Classroom(name="BSIT2B", owner_id=owner.id))


def _student(
    db: Session,
    classroom: Classroom,
    student_id: str,
    full_name: str,
    *,
    program: str | None = None,
    active: bool = True,
) -> Student:
    return StudentRepository(db).create(
        Student(
            classroom_id=classroom.id,
            student_id=student_id,
            full_name=full_name,
            program=program,
            is_active=active,
        )
    )


def _attendance_service(db: Session) -> AttendanceService:
    return AttendanceService(
        AttendanceSessionRepository(db),
        AttendanceRecordRepository(db),
        ClassroomRepository(db),
        StudentRepository(db),
    )


def _history_service(db: Session) -> AttendanceHistoryService:
    return AttendanceHistoryService(
        ClassroomRepository(db),
        AttendanceSessionRepository(db),
        AttendanceRecordRepository(db),
        StudentRepository(db),
    )


def _finalize_day(
    db: Session,
    classroom: Classroom,
    owner: User,
    day: date,
    present: tuple[str, ...] = (),
) -> int:
    """Create, optionally mark, and finalize a session dated ``day``."""
    service = _attendance_service(db)
    session = service.create_session(classroom.id, owner.id)
    session.started_at = datetime(day.year, day.month, day.day, 8, 0, tzinfo=UTC)
    db.commit()
    for business_id in present:
        service.mark_present(session.id, owner.id, student_id=business_id, method="QR")
    service.finalize_session(session.id, owner.id)
    return session.id


def _fix_marked_at(db: Session, session_id: int, business_id: str, marked_at: datetime) -> None:
    record = AttendanceRecordRepository(db).get_by_student_business_id(session_id, business_id)
    assert record is not None
    record.marked_at = marked_at
    db.commit()


def _review(db: Session, session_id: int, owner: User) -> SessionReview:
    return _attendance_service(db).get_session(session_id, owner.id)


def _grid(db: Session, classroom: Classroom, owner: User) -> AttendanceHistoryGrid:
    return _history_service(db).build_grid(classroom.id, owner.id)


def _rows(workbook: Workbook) -> list[list[Any]]:
    """Return the workbook's rows after a save-and-reload round trip."""
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    loaded = load_workbook(buffer)
    worksheet = loaded.active
    assert worksheet is not None
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_session_workbook_contains_the_documented_columns_and_rows(
    db_session: Session,
) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace", program="BSIT")
    _student(db_session, classroom, "25-09760", "Grace Hopper")
    session_id = _finalize_day(
        db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",)
    )
    _fix_marked_at(
        db_session,
        session_id,
        "25-09759",
        datetime(2026, 3, 1, 8, 14, tzinfo=UTC),
    )

    workbook = ExcelService().build_session_workbook(_review(db_session, session_id, owner))

    assert _rows(workbook) == [
        [
            "Student ID",
            "Student Name",
            "Program",
            "Classroom",
            "Date",
            "Status",
            "Marked Time",
            "Method",
        ],
        [
            "25-09759",
            "Ada Lovelace",
            "BSIT",
            "BSIT2B",
            "2026-03-01",
            "PRESENT",
            "08:14 AM",
            "QR",
        ],
        [
            "25-09760",
            "Grace Hopper",
            None,
            "BSIT2B",
            "2026-03-01",
            "ABSENT",
            "-",
            "-",
        ],
    ]


def test_session_export_filename_is_predictable(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    session_id = _finalize_day(
        db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",)
    )

    filename = ExcelService().session_export_filename(_review(db_session, session_id, owner))

    assert filename == "BSIT2B_Attendance_2026-03-01.xlsx"


def test_session_workbook_is_deterministic(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    session_id = _finalize_day(
        db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",)
    )
    review = _review(db_session, session_id, owner)
    service = ExcelService()

    assert _bytes(service.build_session_workbook(review)) == _bytes(
        service.build_session_workbook(review)
    )


def test_session_export_never_modifies_attendance_state(
    db_session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    session_id = _finalize_day(
        db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",)
    )
    review = _review(db_session, session_id, owner)
    commits: list[int] = []
    monkeypatch.setattr(db_session, "commit", lambda: commits.append(1))

    ExcelService().build_session_workbook(review)

    assert commits == []
    record = AttendanceRecordRepository(db_session).get_by_student_business_id(
        session_id, "25-09759"
    )
    assert record is not None
    assert (record.status, record.method) == ("PRESENT", "QR")


def test_history_workbook_transposes_the_grid(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    _student(db_session, classroom, "25-09760", "Grace Hopper")
    _finalize_day(db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",))
    _finalize_day(db_session, classroom, owner, date(2026, 3, 3), present=("25-09760",))

    workbook = ExcelService().build_history_workbook(
        _grid(db_session, classroom, owner), _EXPORTED_AT
    )

    assert _rows(workbook) == [
        ["Student ID", "Student Name", "2026-03-01", "2026-03-03"],
        ["25-09759", "Ada Lovelace", "PRESENT", "ABSENT"],
        ["25-09760", "Grace Hopper", "ABSENT", "PRESENT"],
    ]


def test_history_workbook_leaves_missing_records_blank(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    _finalize_day(db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",))
    _student(db_session, classroom, "25-09761", "Carol Danvers")

    workbook = ExcelService().build_history_workbook(
        _grid(db_session, classroom, owner), _EXPORTED_AT
    )

    rows = _rows(workbook)
    assert rows[0] == ["Student ID", "Student Name", "2026-03-01"]
    assert [value or "" for value in rows[2]] == [
        "25-09761",
        "Carol Danvers",
        "",
    ]


def test_history_workbook_without_finalized_sessions(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")

    workbook = ExcelService().build_history_workbook(
        _grid(db_session, classroom, owner), _EXPORTED_AT
    )

    assert _rows(workbook) == [
        ["Student ID", "Student Name"],
        ["25-09759", "Ada Lovelace"],
    ]


def test_history_workbook_retains_deactivated_students_with_history(
    db_session: Session,
) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    student = _student(db_session, classroom, "25-09759", "Ada Lovelace")
    _finalize_day(db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",))
    student.is_active = False
    db_session.commit()

    workbook = ExcelService().build_history_workbook(
        _grid(db_session, classroom, owner), _EXPORTED_AT
    )

    assert _rows(workbook) == [
        ["Student ID", "Student Name", "2026-03-01"],
        ["25-09759", "Ada Lovelace", "PRESENT"],
    ]


def test_history_export_filename_is_predictable(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    _finalize_day(db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",))

    filename = ExcelService().history_export_filename(
        _grid(db_session, classroom, owner), _EXPORTED_AT
    )

    assert filename == "BSIT2B_Attendance_History_2026-03-02.xlsx"


def test_history_workbook_is_deterministic(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    _student(db_session, classroom, "25-09759", "Ada Lovelace")
    _student(db_session, classroom, "25-09760", "Grace Hopper")
    _finalize_day(db_session, classroom, owner, date(2026, 3, 1), present=("25-09759",))
    _finalize_day(db_session, classroom, owner, date(2026, 3, 3), present=("25-09760",))
    grid = _grid(db_session, classroom, owner)
    service = ExcelService()

    assert _bytes(service.build_history_workbook(grid, _EXPORTED_AT)) == _bytes(
        service.build_history_workbook(grid, _EXPORTED_AT)
    )
