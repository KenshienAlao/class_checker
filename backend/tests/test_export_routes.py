"""Route-level integration tests for the export endpoints."""

from collections.abc import Callable
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any, cast
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from httpx import Response
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.attendance_session import AttendanceSession
from app.models.classroom import Classroom
from app.models.student import Student
from app.models.user import User
from app.repositories.attendance_record_repository import AttendanceRecordRepository
from app.repositories.attendance_session_repository import AttendanceSessionRepository
from app.repositories.classroom_repository import ClassroomRepository
from app.repositories.student_repository import StudentRepository
from app.services.attendance_service import AttendanceService

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _service(db_session: Session) -> AttendanceService:
    """Service wired to the guarded session, for test setup only."""
    return AttendanceService(
        AttendanceSessionRepository(db_session),
        AttendanceRecordRepository(db_session),
        ClassroomRepository(db_session),
        StudentRepository(db_session),
    )


def _create_session(client: TestClient, classroom: Classroom, headers: dict[str, str]) -> Response:
    return cast(
        Response,
        client.post(f"/api/classrooms/{classroom.id}/attendance-sessions", headers=headers),
    )


def _fix_marked_at(
    db_session: Session, session_id: int, business_id: str, marked_at: datetime
) -> None:
    record = AttendanceRecordRepository(db_session).get_by_student_business_id(
        session_id, business_id
    )
    assert record is not None
    record.marked_at = marked_at
    db_session.commit()


def _rows(content: bytes) -> list[list[Any]]:
    """Return the downloaded workbook's rows after a load round trip."""
    workbook = load_workbook(BytesIO(content))
    worksheet = workbook.active
    assert worksheet is not None
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _finalized_session(
    client: TestClient,
    db_session: Session,
    classroom: Classroom,
    headers: dict[str, str],
    present: tuple[str, ...] = (),
) -> int:
    """Create, optionally mark, and finalize a session; return its ID."""
    session_id = cast(int, _create_session(client, classroom, headers).json()["session_id"])
    for business_id in present:
        client.post(
            f"/api/attendance-sessions/{session_id}/scan",
            headers=headers,
            json={"student_id": business_id},
        )
    client.post(f"/api/attendance-sessions/{session_id}/finalize", headers=headers)
    return session_id


def _finalized_session_on(
    client: TestClient,
    db_session: Session,
    classroom: Classroom,
    headers: dict[str, str],
    day: date,
    present: tuple[str, ...] = (),
) -> None:
    """Create a finalized session dated ``day``."""
    session_id = _finalized_session(client, db_session, classroom, headers, present)
    session = db_session.get(AttendanceSession, session_id)
    assert session is not None
    session.started_at = datetime(day.year, day.month, day.day, 8, 0, tzinfo=UTC)
    db_session.commit()


def test_export_requires_a_session(client: TestClient) -> None:
    assert client.get("/api/attendance-sessions/1/export").status_code == 401


def test_export_returns_the_session_workbook(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith", "BSIT")
    make_student(classroom.id, owner.id, "2025-002", "Bob Jones", "BSIT")
    session_id = _finalized_session(client, db_session, classroom, headers, present=("2025-001",))
    _fix_marked_at(
        db_session,
        session_id,
        "2025-001",
        datetime(2026, 3, 1, 8, 14, tzinfo=UTC),
    )
    session = db_session.get(AttendanceSession, session_id)
    assert session is not None
    expected_date = session.started_at.astimezone(UTC).date().isoformat()

    response = client.get(f"/api/attendance-sessions/{session_id}/export", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == _XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="BSIT2B_Attendance_{expected_date}.xlsx"'
    )
    assert _rows(response.content) == [
        [
            "Student ID",
            "Student Name",
            "Program",
            "Classroom",
            "Date",
            "Status",
            "Marked Time",
            "Method",
        ],
        [
            "2025-001",
            "Alice Smith",
            "BSIT",
            "BSIT2B",
            expected_date,
            "PRESENT",
            "08:14 AM",
            "QR",
        ],
        [
            "2025-002",
            "Bob Jones",
            "BSIT",
            "BSIT2B",
            expected_date,
            "ABSENT",
            "-",
            "-",
        ],
    ]


def test_export_returns_404_for_another_owner(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    other = make_user()
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    session_id = _create_session(client, classroom, auth_headers(owner.id)).json()["session_id"]

    response = client.get(
        f"/api/attendance-sessions/{session_id}/export", headers=auth_headers(other.id)
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Attendance session not found"


def test_export_returns_404_for_missing_session(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
) -> None:
    owner = make_user()

    response = client.get(
        "/api/attendance-sessions/999999999/export", headers=auth_headers(owner.id)
    )

    assert response.status_code == 404


def test_export_requires_a_finalized_session(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    session_id = _create_session(client, classroom, headers).json()["session_id"]

    response = client.get(f"/api/attendance-sessions/{session_id}/export", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "Attendance session is not finalized"


def test_export_rejects_a_non_integer_id(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
) -> None:
    owner = make_user()

    response = client.get("/api/attendance-sessions/abc/export", headers=auth_headers(owner.id))

    assert response.status_code == 422


def test_export_never_modifies_attendance_state(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    session_id = _finalized_session(client, db_session, classroom, headers, present=("2025-001",))
    commits: list[int] = []
    monkeypatch.setattr(db_session, "commit", lambda: commits.append(1))

    response = client.get(f"/api/attendance-sessions/{session_id}/export", headers=headers)

    assert response.status_code == 200
    assert commits == []
    session = db_session.get(AttendanceSession, session_id)
    assert session is not None
    assert session.status == "FINALIZED"
    record = AttendanceRecordRepository(db_session).get_by_student_business_id(
        session_id, "2025-001"
    )
    assert record is not None
    assert (record.status, record.method) == ("PRESENT", "QR")


def test_history_export_requires_a_session(client: TestClient) -> None:
    assert client.get("/api/classrooms/1/attendance-history/export").status_code == 401


def test_history_export_returns_the_grid_workbook(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    make_student(classroom.id, owner.id, "2025-002", "Bob Jones")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 3), present=("2025-002",)
    )
    exported_at = datetime.now(UTC)
    expected_date = exported_at.date().isoformat()

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == _XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="BSIT2B_Attendance_History_{expected_date}.xlsx"'
    )
    assert _rows(response.content) == [
        ["Student ID", "Student Name", "2026-03-01", "2026-03-03"],
        ["2025-001", "Alice Smith", "PRESENT", "ABSENT"],
        ["2025-002", "Bob Jones", "ABSENT", "PRESENT"],
    ]


def test_history_export_returns_404_for_another_owner(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    other = make_user()
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export",
        headers=auth_headers(other.id),
    )

    assert response.status_code == 404
    assert response.json()["detail"] == "Classroom not found"


def test_history_export_returns_404_for_missing_classroom(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
) -> None:
    owner = make_user()

    response = client.get(
        "/api/classrooms/999999999/attendance-history/export",
        headers=auth_headers(owner.id),
    )

    assert response.status_code == 404


def test_history_export_with_empty_history(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert _rows(response.content) == [
        ["Student ID", "Student Name"],
        ["2025-001", "Alice Smith"],
    ]


def test_history_export_rejects_a_non_integer_id(
    client: TestClient,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
) -> None:
    owner = make_user()

    response = client.get(
        "/api/classrooms/abc/attendance-history/export", headers=auth_headers(owner.id)
    )

    assert response.status_code == 422


def test_history_export_merges_same_date_sessions_with_present_precedence(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    make_student(classroom.id, owner.id, "2025-002", "Bob Jones")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-002",)
    )

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert _rows(response.content) == [
        ["Student ID", "Student Name", "2026-03-01"],
        ["2025-001", "Alice Smith", "PRESENT"],
        ["2025-002", "Bob Jones", "PRESENT"],
    ]


def test_history_export_retains_inactive_students_with_history(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    student = make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    student.is_active = False
    db_session.commit()

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert _rows(response.content) == [
        ["Student ID", "Student Name", "2026-03-01"],
        ["2025-001", "Alice Smith", "PRESENT"],
    ]


def test_history_export_excludes_inactive_students_without_history(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    departed = make_student(classroom.id, owner.id, "2025-002", "Bob Jones")
    departed.is_active = False
    db_session.commit()
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert _rows(response.content) == [
        ["Student ID", "Student Name", "2026-03-01"],
        ["2025-001", "Alice Smith", "PRESENT"],
    ]


def test_history_export_leaves_blank_cells_for_missing_records(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    make_student(classroom.id, owner.id, "2025-002", "Bob Jones")

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    rows = _rows(response.content)
    assert rows[0] == ["Student ID", "Student Name", "2026-03-01"]
    assert [value or "" for value in rows[2]] == ["2025-002", "Bob Jones", ""]


def test_history_export_excludes_active_sessions(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    _create_session(client, classroom, headers)

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert _rows(response.content) == [
        ["Student ID", "Student Name", "2026-03-01"],
        ["2025-001", "Alice Smith", "PRESENT"],
    ]


def test_history_export_never_modifies_attendance_state(
    client: TestClient,
    db_session: Session,
    make_user: Callable[[], User],
    auth_headers: Callable[[UUID], dict[str, str]],
    make_classroom: Callable[[UUID, str], Classroom],
    make_student: Callable[..., Student],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    owner = make_user()
    headers = auth_headers(owner.id)
    classroom = make_classroom(owner.id, "BSIT2B")
    make_student(classroom.id, owner.id, "2025-001", "Alice Smith")
    _finalized_session_on(
        client, db_session, classroom, headers, date(2026, 3, 1), present=("2025-001",)
    )
    commits: list[int] = []
    monkeypatch.setattr(db_session, "commit", lambda: commits.append(1))

    response = client.get(
        f"/api/classrooms/{classroom.id}/attendance-history/export", headers=headers
    )

    assert response.status_code == 200
    assert commits == []
    session = db_session.scalar(
        select(AttendanceSession).where(
            AttendanceSession.classroom_id == classroom.id,
            AttendanceSession.status == "FINALIZED",
        )
    )
    assert session is not None
    assert session.status == "FINALIZED"
    record = AttendanceRecordRepository(db_session).get_by_student_business_id(
        session.id, "2025-001"
    )
    assert record is not None
    assert (record.status, record.method) == ("PRESENT", "QR")
