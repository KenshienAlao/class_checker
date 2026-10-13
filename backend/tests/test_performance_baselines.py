"""Performance baseline tests backed by the dedicated PostgreSQL test schema."""

from datetime import UTC, datetime, timedelta
from io import BytesIO
from time import perf_counter
from typing import Any
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.engine import Engine
from sqlalchemy.event import listen, remove
from sqlalchemy.orm import Session

from app import models  # noqa: F401 -- register application metadata
from app.core.security import (
    CSRF_COOKIE_NAME,
    SESSION_COOKIE_NAME,
    create_session_token,
    generate_csrf_token,
)
from app.models.attendance_record import AttendanceRecord
from app.models.attendance_session import AttendanceSession
from app.models.classroom import Classroom
from app.models.student import Student
from app.models.user import User
from app.repositories.attendance_record_repository import AttendanceRecordRepository
from app.repositories.attendance_session_repository import AttendanceSessionRepository
from app.repositories.classroom_repository import ClassroomRepository
from app.repositories.student_repository import StudentRepository
from app.services.attendance_service import AttendanceService


class StatementCounter:
    def __init__(self, bind: Engine | Any) -> None:
        self.bind = bind
        self.started_at = 0.0
        self.elapsed_seconds = 0.0
        self.statements: list[str] = []

        def listener(
            conn: Any,
            cursor: Any,
            statement: str,
            parameters: Any,
            context: Any,
            executemany: bool,
        ) -> None:
            self.statements.append(statement)

        self._listener = listener

    @property
    def count(self) -> int:
        return len(self.statements)

    def __enter__(self) -> "StatementCounter":
        self.started_at = perf_counter()
        listen(self.bind, "before_cursor_execute", self._listener)
        return self

    def __exit__(
        self,
        _exc_type: type[BaseException] | None,
        _exc: BaseException | None,
        _tb: Any,
    ) -> Any:
        self.elapsed_seconds = perf_counter() - self.started_at
        remove(self.bind, "before_cursor_execute", self._listener)
        return False


def _owner(db: Session) -> User:
    owner = User(id=uuid4(), name="Owner", email=f"{uuid4()}@example.com", password="n/a")
    db.add(owner)
    db.flush()
    return owner


def _classroom(db: Session, owner: User) -> Classroom:
    return ClassroomRepository(db).create(Classroom(name="Perf C", owner_id=owner.id))


def _students(db: Session, classroom: Classroom, count: int) -> list[Student]:
    students = [
        Student(
            classroom_id=classroom.id,
            student_id=f"s{i:05d}",
            full_name=f"Student {i:05d}",
        )
        for i in range(count)
    ]
    db.add_all(students)
    db.flush()
    return students


def _attendance_service(db: Session) -> AttendanceService:
    return AttendanceService(
        AttendanceSessionRepository(db),
        AttendanceRecordRepository(db),
        ClassroomRepository(db),
        StudentRepository(db),
    )


def _workbook_bytes(row_count: int) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Student ID", "Full Name"])
    for row_number in range(row_count):
        worksheet.append([f"i{row_number:05d}", f"Imported {row_number:05d}"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _auth_headers(user_id: Any) -> dict[str, str]:
    csrf_token = generate_csrf_token()
    return {
        "Cookie": (
            f"{SESSION_COOKIE_NAME}={create_session_token(user_id)};"
            f" {CSRF_COOKIE_NAME}={csrf_token}"
        ),
        "X-CSRF-Token": csrf_token,
    }


@pytest.mark.usefixtures("db_session")
def test_mark_attendance_baseline(db_session: Session) -> None:
    owner = _owner(db_session)
    classroom = _classroom(db_session, owner)
    students = _students(db_session, classroom, 100)
    session = _attendance_service(db_session).create_session(classroom.id, owner.id)
    target_id = students[42].student_id

    counter = StatementCounter(db_session.bind)
    with counter:
        result = _attendance_service(db_session).mark_present(
            session.id, owner.id, student_id=target_id, method="QR"
        )

    records = AttendanceRecordRepository(db_session).list_by_session(session.id)
    assert result.already_marked is False
    assert result.record.status == "PRESENT"
    assert sum(record.status == "PRESENT" for record in records) == 1

    print(
        "PERF baseline mark_attendance records=100 queries=",
        counter.count,
        "elapsed_seconds=",
        counter.elapsed_seconds,
    )


@pytest.mark.usefixtures("db_session")
def test_attendance_finalization_baseline(db_session: Session) -> None:
    for pending_count in (10, 100, 1000):
        owner = _owner(db_session)
        classroom = _classroom(db_session, owner)
        _students(db_session, classroom, pending_count)
        session = _attendance_service(db_session).create_session(classroom.id, owner.id)

        counter = StatementCounter(db_session.bind)
        with counter:
            result = _attendance_service(db_session).finalize_session(session.id, owner.id)

        counts = AttendanceRecordRepository(db_session).count_by_status(session.id)
        assert result.present_count == 0
        assert result.absent_count == pending_count
        assert counts == {"PENDING": 0, "PRESENT": 0, "ABSENT": pending_count}
        print(
            "PERF baseline finalize records=",
            pending_count,
            "queries=",
            counter.count,
            "elapsed_seconds=",
            counter.elapsed_seconds,
        )


@pytest.mark.usefixtures("db_session")
def test_import_students_baseline(
    db_session: Session,
    client: TestClient,
    make_user: Any,
    auth_headers: Any,
    make_classroom: Any,
) -> None:
    owner = make_user()

    for row_count in (1, 100, 1000):
        classroom = make_classroom(owner.id, f"Perf Import {row_count}")
        headers = auth_headers(owner.id)
        content = _workbook_bytes(row_count)
        before = [
            student.student_id
            for student in StudentRepository(db_session).list_active(classroom.id)
        ]

        counter = StatementCounter(db_session.bind)
        with counter:
            response = client.post(
                f"/api/classrooms/{classroom.id}/students/import",
                headers=headers,
                files={
                    "file": (
                        "roster.xlsx",
                        content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

        body = response.json()
        assert response.status_code == 200
        assert body["created"] == row_count
        assert body["skipped"] == 0
        assert body["rejected"] == []
        after = [
            student.student_id
            for student in StudentRepository(db_session).list_active(classroom.id)
        ]
        assert sorted(after) == sorted(before + [f"i{index:05d}" for index in range(row_count)])
        print(
            "PERF baseline import records=",
            row_count,
            "queries=",
            counter.count,
            "elapsed_seconds=",
            counter.elapsed_seconds,
        )


@pytest.mark.usefixtures("db_session")
def test_attendance_export_baseline(
    db_session: Session,
    client: TestClient,
    make_user: Any,
    auth_headers: Any,
    make_classroom: Any,
) -> None:
    owner = make_user()
    classroom = make_classroom(owner.id, "Perf Export")
    classroom_id = classroom.id
    students = [
        Student(classroom_id=classroom_id, student_id=f"e{i:05d}", full_name=f"Export {i:05d}")
        for i in range(1000)
    ]
    db_session.add_all(students)
    db_session.flush()
    session = AttendanceSession(classroom_id=classroom_id, status="ACTIVE")
    db_session.add(session)
    db_session.flush()
    db_session.add_all(
        [
            AttendanceRecord(session_id=session.id, student_id=student.id, status="PENDING")
            for student in students
        ]
    )
    db_session.flush()
    _attendance_service(db_session).finalize_session(session.id, owner.id)
    headers = auth_headers(owner.id)

    counter = StatementCounter(db_session.bind)
    with counter:
        response = client.get(
            f"/api/attendance-sessions/{session.id}/export",
            headers=headers,
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
    worksheet = workbook.active
    assert worksheet is not None
    rows = list(worksheet.iter_rows(values_only=True))
    assert len(rows) == 1001
    assert rows[0][0] == "Student ID"
    print(
        "PERF baseline export records=1000 queries=",
        counter.count,
        "elapsed_seconds=",
        counter.elapsed_seconds,
    )


@pytest.mark.usefixtures("db_session")
def test_attendance_history_baseline(
    db_session: Session,
    client: TestClient,
    make_user: Any,
    auth_headers: Any,
    make_classroom: Any,
) -> None:
    owner = make_user()
    classroom = make_classroom(owner.id, "Perf History")
    students = [
        Student(classroom_id=classroom.id, student_id=f"h{i:05d}", full_name=f"History {i:05d}")
        for i in range(25)
    ]
    db_session.add_all(students)
    db_session.flush()

    for session_index in range(100):
        session = AttendanceSession(
            classroom_id=classroom.id,
            status="ACTIVE",
            started_at=datetime(2026, 1, 1, tzinfo=UTC) + timedelta(days=session_index),
        )
        db_session.add(session)
        db_session.flush()
        db_session.add_all(
            [
                AttendanceRecord(session_id=session.id, student_id=student.id, status="PENDING")
                for student in students
            ]
        )
        db_session.flush()
        _attendance_service(db_session).finalize_session(session.id, owner.id)

    headers = auth_headers(owner.id)
    counter = StatementCounter(db_session.bind)
    with counter:
        response = client.get(
            f"/api/classrooms/{classroom.id}/attendance-history",
            headers=headers,
        )

    body = response.json()
    assert response.status_code == 200
    assert len(body["dates"]) == 100
    assert len(body["rows"]) == 25
    for row in body["rows"]:
        assert len(row["statuses"]) == 100
    print(
        "PERF baseline history sessions=100 students=25 queries=",
        counter.count,
        "elapsed_seconds=",
        counter.elapsed_seconds,
    )
